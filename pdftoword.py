import os
import io
import tempfile
import tkinter as tk
from tkinter import filedialog, messagebox

import fitz  # PyMuPDF
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from tqdm import tqdm


def convert_pdf(pdf_path):
    try:
        pdf = fitz.open(pdf_path)

        word = Document()

        style = word.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        temp_dir = tempfile.mkdtemp()

        total_pages = len(pdf)

        for page_no in tqdm(
                range(total_pages),
                desc=os.path.basename(pdf_path),
                leave=False):

            page = pdf.load_page(page_no)

            # Extract text
            text = page.get_text("text")
            if text.strip():
                word.add_paragraph(text)

            # Extract images
            images = page.get_images(full=True)

            for img_no, img in enumerate(images):

                try:
                    xref = img[0]
                    base_image = pdf.extract_image(xref)

                    image_bytes = base_image["image"]
                    ext = base_image["ext"]

                    image_path = os.path.join(
                        temp_dir,
                        f"{page_no}_{img_no}.{ext}"
                    )

                    with open(image_path, "wb") as f:
                        f.write(image_bytes)

                    try:
                        im = Image.open(image_path)

                        if im.width > 1000:
                            word.add_picture(image_path, width=Inches(6))
                        else:
                            word.add_picture(image_path)

                    except:
                        pass

                except:
                    pass

            if page_no != total_pages - 1:
                word.add_page_break()

        folder = os.path.dirname(pdf_path)
        filename = os.path.splitext(os.path.basename(pdf_path))[0]

        output = os.path.join(
            folder,
            filename + "_converted.docx"
        )

        word.save(output)

        pdf.close()

        return True

    except Exception as e:
        print(f"\nFailed : {pdf_path}")
        print(e)
        return False


def main():

    root = tk.Tk()
    root.withdraw()

    pdf_files = filedialog.askopenfilenames(
        title="Select PDF File(s)",
        filetypes=[("PDF Files", "*.pdf")]
    )

    if not pdf_files:
        return

    success = 0

    print()

    for pdf in pdf_files:

        print(f"Converting : {os.path.basename(pdf)}")

        if convert_pdf(pdf):
            success += 1

    messagebox.showinfo(
        "Completed",
        f"{success} of {len(pdf_files)} PDF(s) converted successfully."
    )


if __name__ == "__main__":
    main()